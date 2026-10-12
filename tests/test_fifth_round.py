"""Verkliga processgränser, köer, avbrott och atomiska generationsbyten."""
import asyncio
import ctypes
import json
import os
from pathlib import Path
import sys
import time

import pytest

import crawler_runtime as runtime
from test_crawler import Site, config, make_pdf, page, published, run, run_crawl, serve, uwc


def worker_script(tmp_path, default='success'):
    script = tmp_path / ('worker_' + default + '.py')
    script.write_text('''import json, os, pathlib, subprocess, sys, time
if sys.stdin.buffer.read(1) != b'G':
    sys.exit(2)
request = json.loads(pathlib.Path(sys.argv[1]).read_text(encoding='utf-8'))
options = request['options']
action = options.get('action', DEFAULT_ACTION)
if action == 'crash':
    os._exit(77)
if action == 'child':
    child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])
    pathlib.Path(options['pid_file']).write_text(str(child.pid))
if action in ('timeout', 'child'):
    time.sleep(60)
time.sleep(options.get('delay', 0))
folder = pathlib.Path(request['output']) / 'texter'
folder.mkdir(parents=True)
path = folder / options.get('md_filename', 'output_doc.md')
path.write_text('# Processdokument\\n\\n**Dokument-URL:** ' + request['url'] +
                '\\n\\n---\\n\\nTEXT FRAN EGEN PROCESS ' + str(os.getpid()) + '\\n\\n---\\n', encoding='utf-8')
if action == 'escape':
    path = pathlib.Path(options['escape_path'])
result = {'path': str(path), 'ocr': False, 'error': ''}
pathlib.Path(sys.argv[2]).write_text(json.dumps(result), encoding='utf-8')
'''.replace('DEFAULT_ACTION', repr(default)), encoding='utf-8')
    return str(script)


def request_options(**kwargs):
    return {'md_filename': 'output_doc.md', **kwargs}


def process_is_running(pid):
    if os.name == 'nt':
        from ctypes import wintypes as wt
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.OpenProcess.argtypes = [wt.DWORD, wt.BOOL, wt.DWORD]
        kernel.OpenProcess.restype = wt.HANDLE
        kernel.GetExitCodeProcess.argtypes = [wt.HANDLE, ctypes.POINTER(wt.DWORD)]
        kernel.CloseHandle.argtypes = [wt.HANDLE]
        handle = kernel.OpenProcess(0x1000, False, pid)
        if not handle:
            return False
        code = wt.DWORD()
        try:
            return bool(kernel.GetExitCodeProcess(handle, ctypes.byref(code))) and code.value == 259
        finally:
            kernel.CloseHandle(handle)
    try:
        stat = Path(f'/proc/{pid}/stat')
        if stat.exists() and stat.read_text().split()[2] == 'Z':
            return False
        os.kill(pid, 0)
        return True
    except ProcessLookupError:
        return False


@pytest.mark.parametrize('failure', ['crash', 'timeout', 'escape'])
def test_failed_worker_does_not_poison_later_jobs_or_leave_processes(tmp_path, failure):
    async def go():
        source = tmp_path / 'input.pdf'
        source.write_bytes(b'%PDF-TEST')
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=0.6,
                                                  worker_script=worker_script(tmp_path))
        try:
            first = await runner.run(str(source), 'https://test.se/input.pdf', {},
                                     request_options(action=failure, escape_path=str(source)),
                                     'output_doc.md', 10000)
            assert first['text'] is None and first['error']
            assert not runner.active
            assert not list((tmp_path / 'jobs').iterdir())
            second = await runner.run(str(source), 'https://test.se/input.pdf', {},
                                      request_options(), 'output_doc.md', 10000)
            assert 'TEXT FRAN EGEN PROCESS' in second['text']
            assert str(os.getpid()) != second['text'].split('TEXT FRAN EGEN PROCESS ')[1].split()[0]
        finally:
            await runner.close()
    run(go())


def test_worker_timeout_kills_descendant_processes(tmp_path):
    async def go():
        source = tmp_path / 'input.pdf'
        source.write_bytes(b'%PDF-TEST')
        pid_file = tmp_path / 'child.pid'
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=0.8,
                                                  worker_script=worker_script(tmp_path))
        try:
            result = await runner.run(str(source), 'https://test.se/input.pdf', {},
                                      request_options(action='child', pid_file=str(pid_file)),
                                      'output_doc.md', 10000)
            assert result['text'] is None and 'tidsgränsen' in result['error']
            pid = int(pid_file.read_text())
            deadline = time.monotonic() + 3
            while process_is_running(pid) and time.monotonic() < deadline:
                await asyncio.sleep(0.05)
            assert not process_is_running(pid)
        finally:
            await runner.close()
    run(go())


def test_conversion_process_concurrency_is_bounded(tmp_path):
    async def go():
        source = tmp_path / 'input.pdf'
        source.write_bytes(b'%PDF-TEST')
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), workers=2, timeout=3,
                                                  worker_script=worker_script(tmp_path))
        tasks = [asyncio.create_task(runner.run(str(source), 'https://test.se/' + str(i), {},
                                                request_options(delay=0.15), 'output_doc.md', 10000))
                 for i in range(6)]
        peak = 0
        try:
            while not all(task.done() for task in tasks):
                peak = max(peak, len(runner.active))
                await asyncio.sleep(0.01)
            assert peak == 2
            assert all(result['text'] for result in await asyncio.gather(*tasks))
            assert not runner.active
        finally:
            await runner.close()
    run(go())


@pytest.mark.parametrize('kind', ['pdf', 'docx', 'pptx', 'xlsx'])
def test_real_worker_extracts_and_cleans_documents_in_another_process(tmp_path, monkeypatch, kind):
    source = tmp_path / ('report.' + kind)
    text = 'Information i dokumentet, kontakta person@example.org for flera uppgifter.'
    if kind == 'pdf':
        source.write_bytes(make_pdf(text))
    elif kind == 'docx':
        from docx import Document
        doc = Document()
        doc.add_paragraph(text)
        doc.save(source)
    elif kind == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = text
        prs.save(source)
    else:
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(['Information'])
        wb.active.append([text])
        wb.save(source)
        wb.close()
    pids = []
    original = runtime.asyncio.create_subprocess_exec
    async def observed(*args, **kwargs):
        process = await original(*args, **kwargs)
        pids.append(process.pid)
        return process
    monkeypatch.setattr(runtime.asyncio, 'create_subprocess_exec', observed)
    async def go():
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=20)
        url = 'https://test.se/report.' + kind
        try:
            result = await runner.run(str(source), url, {}, {'remove_email': True, 'ocr': False},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000,
                                      expected_hash=uwc.file_sha256(str(source)))
            assert result['text'], result['error']
            assert 'Information' in result['text']
            assert '[E-POST]' in result['text'] and 'person@example.org' not in result['text']
            assert pids and all(pid != os.getpid() for pid in pids)
            assert not list((tmp_path / 'jobs').iterdir())
        finally:
            await runner.close()
    run(go())


def test_changed_input_version_is_not_converted(tmp_path):
    async def go():
        source = tmp_path / 'report.pdf'
        source.write_bytes(make_pdf('Dokumentets text och annat innehall.'))
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=20)
        url = 'https://test.se/report.pdf'
        try:
            result = await runner.run(str(source), url, {}, {},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000,
                                      expected_hash='0' * 64)
            assert result['text'] is None
            assert 'Originalversionen' in result['error']
        finally:
            await runner.close()
    run(go())


def test_internal_worker_entrypoint_uses_loaded_program_module(tmp_path, monkeypatch):
    async def go():
        source = tmp_path / 'report.pdf'
        source.write_bytes(make_pdf('DOKUMENT VIA INTERN WORKER-INGANG'))
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=20)
        script = Path(uwc.__file__)
        monkeypatch.setattr(runner, 'command', lambda job, result: [sys.executable, str(script),
                                                                   '--conversion-worker', job, result])
        url = 'https://test.se/report.pdf'
        try:
            result = await runner.run(str(source), url, {}, {'ocr': False},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000)
            assert result['text'], result['error']
            assert 'INTERN WORKER-INGANG' in result['text']
        finally:
            await runner.close()
    run(go())


@pytest.mark.skipif(os.name != 'nt', reason='pythonw finns endast på Windows')
def test_windowless_python_worker_does_not_require_sys_stdin(tmp_path, monkeypatch):
    executable = Path(sys.executable).with_name('pythonw.exe')
    if not executable.is_file():
        pytest.skip('pythonw.exe saknas')
    async def go():
        source = tmp_path / 'report.pdf'
        source.write_bytes(make_pdf('DOKUMENT MED FONSTERLOS PYTHON'))
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'), timeout=20)
        monkeypatch.setattr(runner, 'command', lambda job, result: [str(executable), '-u',
                                                                   runner.worker_script, job, result])
        url = 'https://test.se/report.pdf'
        try:
            result = await runner.run(str(source), url, {}, {'ocr': False},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000)
            assert result['text'], result['error']
            assert 'FONSTERLOS PYTHON' in result['text']
        finally:
            await runner.close()
    run(go())


def test_gui_templates_keep_conversion_options_and_clear_them_when_unselected():
    from types import SimpleNamespace
    gui = SimpleNamespace(lang='sv', templates={'Dokument': {'conversion_workers': 1, 'conversion_timeout': 300}},
                          texts={'template_none': {'sv': 'Ingen mall'}, 'template_loaded': {'sv': 'Laddad {}'}},
                          _log_to_gui=lambda *args: None)
    uwc.AppGUI._apply_template(gui, 'Dokument')
    assert gui._template_conversion_options == {'conversion_workers': 1, 'conversion_timeout': 300}
    uwc.AppGUI._apply_template(gui, 'Ingen mall')
    assert gui._template_conversion_options == {}


def test_document_conversion_does_not_occupy_a_crawl_slot(tmp_path, monkeypatch):
    async def go():
        gate, started, visited = asyncio.Event(), asyncio.Event(), asyncio.Event()
        async def slow_result(self, filepath, url, ref, *args, **kwargs):
            started.set()
            await gate.wait()
            return {'text': '# Rapport\n\n**Dokument-URL:** ' + url +
                    '\n\n---\n\nKonverterat dokument med tillrackligt mycket text.\n\n---\n',
                    'ocr': False, 'error': ''}
        monkeypatch.setattr(uwc.ConversionProcessRunner, 'run', slow_result)
        class ObservedSite(Site):
            async def handle(self, request):
                if request.path == '/later':
                    visited.set()
                return await super().handle(request)
        site = ObservedSite()
        site.pages['/'] = (page('Start', 'Starttext', [('/report.pdf', 'Rapport'), ('/later', 'Senare')]), None)
        site.pages['/later'] = (page('Senare', 'SIDAN HAR HAMTATS'), None)
        site.files['/report.pdf'] = (make_pdf('Dokument'), None)
        server, base = await serve(site)
        c = uwc.AsyncWebCrawler(config(base, tmp_path, concurrency=1, conversion_workers=1, find_sitemap=False))
        task = asyncio.create_task(c.crawl())
        try:
            await asyncio.wait_for(started.wait(), 3)
            await asyncio.wait_for(visited.wait(), 3)
            assert not gate.is_set(), 'Sidan måste hämtas medan dokumentkonverteringen väntar'
            gate.set()
            await asyncio.wait_for(task, 10)
            assert 'Konverterat dokument' in (published(tmp_path) / 'chunks.jsonl').read_text(encoding='utf-8')
            assert c._conversion_pending == 0
        finally:
            gate.set()
            if not task.done():
                c.stop()
                await task
            await server.cleanup()
    run(go())


def test_stop_cancels_real_workers_and_bounded_queue_without_publishing(tmp_path):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'URSPRUNGLIG PUBLICERING'), 's1')
        server, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, find_sitemap=False)
            before = (tmp_path / 'current.json').read_bytes()
            site.pages['/'] = (page('Start', 'NY VERSION', [(f'/r{i}.pdf', str(i)) for i in range(4)]), 's2')
            for i in range(4):
                site.files[f'/r{i}.pdf'] = (make_pdf('Ny rapport'), None)
            c = uwc.AsyncWebCrawler(config(base, tmp_path, conversion_workers=1, conversion_queue_size=1,
                                           concurrency=4, find_sitemap=False, conversion_timeout=60))
            c._conversion_runner.worker_script = worker_script(tmp_path, default='timeout')
            task = asyncio.create_task(c.crawl())
            async with asyncio.timeout(5):
                while not c._conversion_runner.active or c._conversion_queue.qsize() < 1:
                    await asyncio.sleep(0.02)
            assert c._conversion_queue.maxsize == c._conversion_queue.qsize() == 1
            c.stop()
            await asyncio.wait_for(task, 5)
            assert not c._conversion_runner.active
            assert not c._conversion_tasks and c._conversion_pending == 0
            assert (tmp_path / 'current.json').read_bytes() == before
            assert 'URSPRUNGLIG PUBLICERING' in (published(tmp_path) / 'chunks.jsonl').read_text(encoding='utf-8')
            assert not list((tmp_path / '.working').iterdir())
        finally:
            await server.cleanup()
    run(go())


@pytest.mark.parametrize('artifact', ['index', 'manifest', 'jsonl', 'pointer'])
def test_publication_failures_preserve_entire_previous_generation(tmp_path, monkeypatch, artifact):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'GAMMAL GENERATION'), 's1')
        server, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, find_sitemap=False)
            pointer_before = (tmp_path / 'current.json').read_bytes()
            old = published(tmp_path)
            data_before = {name: (old / name).read_bytes() for name in ('chunks.jsonl', 'index.csv', 'changes.jsonl')}
            site.pages['/'] = (page('Ny titel', 'NY GENERATION'), 's2')
            with monkeypatch.context() as patch:
                if artifact == 'pointer':
                    replace = runtime.os.replace
                    def failed_replace(source, target):
                        if str(target).endswith('current.json'):
                            raise OSError('Avbrott vid pointer-byte')
                        return replace(source, target)
                    patch.setattr(runtime.os, 'replace', failed_replace)
                else:
                    target = {'index': '_generate_index', 'manifest': '_generate_manifest', 'jsonl': '_write_chunks_jsonl'}[artifact]
                    if artifact == 'jsonl':
                        def failed(*args):
                            raise OSError('Exportfel')
                    else:
                        async def failed(*args):
                            raise OSError('Exportfel')
                    patch.setattr(uwc.AsyncWebCrawler, target, failed)
                result = await run_crawl(base, tmp_path, find_sitemap=False)
            assert result.delivery_status() == 'failed'
            assert (tmp_path / 'current.json').read_bytes() == pointer_before
            assert published(tmp_path) == old
            assert all((old / name).read_bytes() == value for name, value in data_before.items())
            assert 'GAMMAL GENERATION' in (tmp_path / 'chunks.jsonl').read_text(encoding='utf-8')
            assert not list((tmp_path / '.working').iterdir())
            restored = await run_crawl(base, tmp_path, find_sitemap=False)
            assert restored.delivery_status() == 'success'
            assert 'NY GENERATION' in (published(tmp_path) / 'chunks.jsonl').read_text(encoding='utf-8')
        finally:
            await server.cleanup()
    run(go())


def test_reader_pins_old_generation_while_new_generation_is_published(tmp_path):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'VERSION ETT'), 's1')
        server, base = await serve(site)
        try:
            first = await run_crawl(base, tmp_path, find_sitemap=False)
            old = published(tmp_path)
            site.pages['/'] = (page('Start', 'VERSION TVA'), 's2')
            second = await run_crawl(base, tmp_path, find_sitemap=False)
            new = published(tmp_path)
            assert first.crawl_session_id != second.crawl_session_id
            assert old != new
            assert 'VERSION ETT' in (old / 'chunks.jsonl').read_text(encoding='utf-8')
            assert 'VERSION TVA' in (new / 'chunks.jsonl').read_text(encoding='utf-8')
            pointer = json.loads((tmp_path / 'current.json').read_text())
            report = json.loads((new / 'crawl_report.json').read_text())
            assert pointer['run'] == report['run'] == second.crawl_session_id
            assert all((new / name).is_file() for name in pointer['artifacts'])
            assert all(not name.endswith(('.tmp', '.part', '-wal', '-shm')) for name in pointer['artifacts'])
            assert pointer['corpus'] == 'chunks.jsonl'
            assert all(not name.startswith('texter_utgangna/') and not name.endswith('_cache.db')
                       for name in pointer['artifacts'])
        finally:
            await server.cleanup()
    run(go())


def test_modifying_legacy_copy_cannot_change_published_generation(tmp_path):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'PUBLICERAD TEXT'), 's1')
        server, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, find_sitemap=False)
            old = published(tmp_path)
            name = next((old / 'texter').glob('*.md')).name
            (tmp_path / 'texter' / name).write_text('REDIGERAD KOMPATIBILITETSKOPIA')
            assert 'PUBLICERAD TEXT' in (old / 'texter' / name).read_text(encoding='utf-8')
            await run_crawl(base, tmp_path, find_sitemap=False)
            assert 'PUBLICERAD TEXT' in (published(tmp_path) / 'chunks.jsonl').read_text(encoding='utf-8')
        finally:
            await server.cleanup()
    run(go())


def test_disabled_jsonl_does_not_publish_an_old_export(tmp_path):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'VERSION ETT'), 's1')
        server, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, find_sitemap=False)
            assert (published(tmp_path) / 'chunks.jsonl').exists()
            site.pages['/'] = (page('Start', 'VERSION TVA'), 's2')
            await run_crawl(base, tmp_path, find_sitemap=False, export_jsonl=False)
            assert not (published(tmp_path) / 'chunks.jsonl').exists()
            assert not (tmp_path / 'chunks.jsonl').exists()
        finally:
            await server.cleanup()
    run(go())


def test_output_lock_and_invalid_pointer_do_not_modify_other_data(tmp_path):
    root = tmp_path / 'output'
    first = runtime.PublicationWorkspace(str(root), 'cache.db', 'a')
    first.begin()
    try:
        other = runtime.PublicationWorkspace(str(root), 'cache.db', 'b')
        with pytest.raises(RuntimeError, match='annan crawl'):
            other.begin()
    finally:
        first.abort()
    (root / 'current.json').write_text(json.dumps({'generation': '../outside'}))
    other = runtime.PublicationWorkspace(str(root), 'cache.db', 'b')
    with pytest.raises(ValueError, match='ogiltig generation'):
        other.begin()
    assert other.lock is None
    assert not list((root / '.working').iterdir())


def test_recursive_cleanup_rejects_targets_outside_its_owner(tmp_path):
    root = tmp_path / 'root'
    root.mkdir()
    other = tmp_path / 'outside'
    other.mkdir()
    (other / 'keep.txt').write_text('Bevara')
    with pytest.raises(ValueError):
        runtime.remove_tree(str(root), str(other))
    with pytest.raises(ValueError):
        runtime.remove_tree(str(root), str(root))
    assert (other / 'keep.txt').exists()


def test_pointer_replace_retries_a_transient_reader_lock(tmp_path, monkeypatch):
    target = tmp_path / 'current.json'
    target.write_text('{"old": true}')
    original = runtime.os.replace
    calls = []
    def locked_once(source, destination):
        calls.append(destination)
        if len(calls) == 1:
            raise PermissionError('Läsaren håller filen öppen')
        return original(source, destination)
    monkeypatch.setattr(runtime.os, 'replace', locked_once)
    runtime.write_json(str(target), {'new': True}, retry_replace=True)
    assert len(calls) == 2
    assert json.loads(target.read_text()) == {'new': True}


def test_stale_workspaces_are_removed_after_a_crashed_writer(tmp_path):
    root = tmp_path / 'output'
    stale = root / '.working' / ('a' * 32)
    stale.mkdir(parents=True)
    (stale / 'half-written.part').write_text('Ofärdigt')
    workspace = runtime.PublicationWorkspace(str(root), 'cache.db', 'restart')
    workspace.begin()
    try:
        assert not stale.exists()
        assert Path(workspace.workdir).exists()
    finally:
        workspace.abort()


def test_corrupt_published_original_forces_a_body_fetch_instead_of_304(tmp_path):
    async def go():
        site = Site()
        site.pages['/'] = (page('Start', 'Starttext', [('/report.pdf', 'Rapport')]), 's1')
        site.files['/report.pdf'] = (make_pdf('KORREKT ORIGINAL'), 'd1')
        server, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, find_sitemap=False)
            next((published(tmp_path) / 'dokument').glob('*.pdf')).write_bytes(make_pdf('KORRUPT VERSION'))
            site.hits.clear()
            await run_crawl(base, tmp_path, find_sitemap=False)
            assert ('/report.pdf', False, 'GET') in site.hits
            text = (published(tmp_path) / 'chunks.jsonl').read_text(encoding='utf-8')
            assert 'KORREKT ORIGINAL' in text and 'KORRUPT VERSION' not in text
        finally:
            await server.cleanup()
    run(go())


@pytest.mark.parametrize('key,value', [('conversion_workers', 0), ('conversion_queue_size', 0),
                                     ('conversion_timeout', 0), ('conversion_timeout', float('inf'))])
def test_conversion_configuration_is_validated(key, value):
    errors, _ = uwc.validate_site_config({'start_url': 'https://test.se', key: value})
    assert any(key in error for error in errors)


def test_real_libreoffice_conversion_when_installed(tmp_path):
    converter = uwc.DocumentConverter(str(tmp_path))
    if not converter._soffice:
        pytest.skip('LibreOffice är inte installerat i testmiljön')
    async def go():
        source = tmp_path / 'legacy.rtf'
        source.write_text(r'{\rtf1\ansi Verklig LibreOffice-konvertering av riktlinjer och avgifter.}')
        url = 'https://test.se/legacy.rtf'
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'))
        try:
            result = await runner.run(str(source), url, {}, {},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000)
            assert result['text'], result['error']
            assert 'LibreOffice-konvertering' in result['text']
        finally:
            await runner.close()
    run(go())


def test_real_ocr_worker_when_language_data_is_installed(tmp_path):
    converter = uwc.DocumentConverter(str(tmp_path))
    if not converter.ocr_available():
        pytest.skip('Tesseracts språkdata saknas i testmiljön')
    async def go():
        doc = uwc.pymupdf.open()
        p = doc.new_page()
        p.insert_text((50, 100), 'Municipal parking fee 1250 kronor', fontsize=24)
        image = p.get_pixmap(matrix=uwc.pymupdf.Matrix(2, 2))
        scanned = uwc.pymupdf.open()
        target = scanned.new_page()
        target.insert_image(target.rect, pixmap=image)
        source = tmp_path / 'scanned.pdf'
        scanned.save(source)
        scanned.close()
        doc.close()
        runner = runtime.ConversionProcessRunner(str(tmp_path / 'jobs'))
        url = 'https://test.se/scanned.pdf'
        try:
            result = await runner.run(str(source), url, {}, {'ocr': True, 'ocr_languages': 'eng'},
                                      uwc.DocumentConverter.md_filename(url, source.name), 1000000)
            assert result['text'] and result['ocr'], result['error']
            assert '1250' in result['text']
        finally:
            await runner.close()
    run(go())
