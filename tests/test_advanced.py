"""Tester för Playwright-vägar, inloggning, OCR/LibreOffice-koppling och innehållskvalitet.

Playwright-testerna kräver Chromium (`playwright install chromium`) och hoppas annars över.
Alla tester går mot en lokal testserver.
"""
import asyncio
import json
import os
import queue
import shutil
import sys

import pytest
from aiohttp import web

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_crawler import (Site, LOREM, config, make_pdf, page, read_changes,  # noqa: E402
                          run, run_crawl, serve, texts, uwc)


def chromium_available() -> bool:
    try:
        async def probe():
            from playwright.async_api import async_playwright
            async with async_playwright() as p:
                b = await p.chromium.launch(headless=True)
                await b.close()
        asyncio.run(probe())
        return True
    except Exception:
        return False


needs_chromium = pytest.mark.skipif(not chromium_available(),
                                    reason="Playwright Chromium saknas")


def all_md(out) -> str:
    return "".join(open(os.path.join(out, "texter", t), encoding="utf-8").read()
                   for t in texts(out))


# ───────────────────────── Playwright ─────────────────────────
@needs_chromium
def test_javascript_page_is_rendered_with_playwright(tmp_path):
    spa = ("<html><head><title>SPA</title></head><body><div id='root'></div><script>"
           "document.getElementById('root').innerHTML = '<h1>Renderad rubrik</h1><p>' + "
           "'Renderat innehåll från JavaScript. '.repeat(10) + '</p>';</script></body></html>")

    async def go():
        site = Site()
        site.pages["/"] = (spa, None)
        runner, base = await serve(site)
        try:
            return await run_crawl(base, tmp_path)
        finally:
            await runner.cleanup()
    c = run(go())
    assert c.stats.playwright_fallbacks == 1
    assert "Renderat innehåll från JavaScript" in all_md(tmp_path)


class AuthSite(Site):
    """Sajt där allt utom inloggningssidan kräver cookien sess=<token>."""

    def __init__(self):
        super().__init__()
        self.token = "t1"

    async def handle(self, request):
        has = request.cookies.get("sess") == self.token
        if request.path == "/" and not has:
            # "Inloggningssida" som sätter cookien via JS (simulerar att användaren loggat in)
            html = ("<html><head><title>Logga in</title></head><body>"
                    "<form id='loginform'><input type='password'></form>"
                    f"<script>document.cookie = 'sess={self.token}; path=/';</script>"
                    + "<p>Logga in för att fortsätta.</p>" * 20 + "</body></html>")
            return web.Response(text=html, content_type="text/html")
        if request.path != "/" and not has and request.path not in ("/robots.txt", "/sitemap.xml"):
            return web.Response(status=302, headers={"Location": "/"})
        return await super().handle(request)


def login_cfg(base, out, cookie_file, **kw):
    return config(base, out, headless_mode="login_then_headless", cookie_file=str(cookie_file),
                  login_browser_headless=True, respect_robots=False, find_sitemap=False,
                  download_docs=False, convert_docs_to_md=False, **kw)


async def crawl_with_auto_login(cfg):
    """Kör en crawl med GUI-kö och 'klickar OK' så fort inloggningsdialogen efterfrågas."""
    q = queue.Queue()
    crawler = uwc.AsyncWebCrawler(cfg, q)

    async def user():
        while True:
            try:
                kind, _ = q.get_nowait()
            except queue.Empty:
                await asyncio.sleep(0.05)
                continue
            if kind == "login_wait":
                crawler.login_event.set()
                return
    await asyncio.wait_for(asyncio.gather(crawler.crawl(), user()), timeout=90)
    return crawler


@needs_chromium
def test_interactive_login_cookie_file_reuse_and_expiry(tmp_path):
    cookie_file = tmp_path / "cookies.json"

    async def go():
        site = AuthSite()
        site.pages["/"] = (page("Intranät", "start", [("/a", "A")]), None)
        site.pages["/a"] = (page("Sida A", "a-innehåll"), None)
        runner, base = await serve(site)
        results = {}
        try:
            # 1. Interaktiv inloggning (GUI-läge): användaren "klickar OK" → cookies tas över
            out1 = tmp_path / "run1"
            results["c1"] = await crawl_with_auto_login(login_cfg(base, out1, cookie_file))
            results["out1"] = out1

            # 2. Serverläge (ingen GUI): återanvänder cookie-filen, ingen webbläsare behövs
            out2 = tmp_path / "run2"
            c2 = uwc.AsyncWebCrawler(login_cfg(base, out2, cookie_file))
            await asyncio.wait_for(c2.crawl(), timeout=60)
            results["c2"] = c2
            results["out2"] = out2

            # 3. Sessionen har gått ut (servern byter token) i serverläge → tydligt fel
            site.token = "t2"
            out3 = tmp_path / "run3"
            c3 = uwc.AsyncWebCrawler(login_cfg(base, out3, cookie_file))
            await asyncio.wait_for(c3.crawl(), timeout=60)
            results["c3"] = c3
            results["out3"] = out3

            # 4. Samma sak i GUI-läge → loggar in på nytt automatiskt
            out4 = tmp_path / "run4"
            results["c4"] = await crawl_with_auto_login(login_cfg(base, out4, cookie_file))
            results["out4"] = out4
        finally:
            await runner.cleanup()
        return results
    r = run(go())

    assert r["c1"].stats.pages_visited == 2, r["c1"].fatal_error
    assert "Sida A" in all_md(r["out1"])
    assert cookie_file.exists()
    assert any(c["name"] == "sess" for c in json.loads(cookie_file.read_text()))

    assert r["c2"].stats.pages_visited == 2, r["c2"].fatal_error
    assert "Sida A" in all_md(r["out2"])

    assert r["c3"].fatal_error and "cookie" in r["c3"].fatal_error.lower()
    assert r["c3"].stats.pages_visited == 0
    assert texts(r["out3"]) == []

    assert r["c4"].stats.pages_visited == 2, r["c4"].fatal_error
    assert "Sida A" in all_md(r["out4"])


def test_server_mode_without_cookie_file_fails_clearly(tmp_path):
    async def go():
        site = AuthSite()
        site.pages["/"] = (page("Intranät", "start"), None)
        runner, base = await serve(site)
        try:
            c = uwc.AsyncWebCrawler(login_cfg(base, tmp_path, tmp_path / "saknas.json"))
            await asyncio.wait_for(c.crawl(), timeout=60)
            return c
        finally:
            await runner.cleanup()
    c = run(go())
    assert c.fatal_error and "cookie_file" in c.fatal_error
    assert texts(tmp_path) == []


@needs_chromium
def test_document_download_falls_back_to_playwright(tmp_path):
    class BrowserOnlySite(Site):
        async def handle(self, request):
            if request.path.startswith("/d/") and "Sec-Fetch-Mode" not in request.headers:
                return web.Response(status=403, text="endast webbläsare")
            if request.path.startswith("/d/"):
                data, _ = self.files[request.path]
                return web.Response(body=data, content_type="application/pdf", headers={
                    "Content-Disposition": 'attachment; filename="rapport.pdf"'})
            return await super().handle(request)

    async def go():
        site = BrowserOnlySite()
        site.pages["/"] = (page("Start", "start", [("/d/rapport.pdf", "Rapport för 2025")]), None)
        site.files["/d/rapport.pdf"] = (make_pdf("Rapporten som bara webbläsare får hämta"), None)
        runner, base = await serve(site)
        try:
            return await run_crawl(base, tmp_path)
        finally:
            await runner.cleanup()
    c = run(go())
    assert c.stats.documents_downloaded == 1
    docs = os.listdir(tmp_path / "dokument")
    assert len(docs) == 1 and docs[0].endswith(".pdf") and not docs[0].endswith(".part")
    assert "bara webbläsare" in all_md(tmp_path)


# ───────────────────────── LibreOffice / OCR (kopplingen) ─────────────────────────
def test_legacy_format_conversion_plumbing(tmp_path, monkeypatch):
    """Riktiga LibreOffice finns inte här — vi testar kopplingen med en attrapp."""
    from docx import Document
    template = tmp_path / "fran_soffice.docx"
    d = Document()
    d.add_heading("Gammal rubrik", 1)
    d.add_paragraph("Innehåll som LibreOffice konverterat från ett gammalt .doc-dokument.")
    d.save(template)

    calls = []

    def fake_run(cmd, **kw):
        calls.append(cmd)
        outdir = cmd[cmd.index("--outdir") + 1]
        base = os.path.splitext(os.path.basename(cmd[-1]))[0]
        shutil.copy(template, os.path.join(outdir, base + ".docx"))

    monkeypatch.setattr(uwc.subprocess, "run", fake_run)
    conv = uwc.DocumentConverter(str(tmp_path / "out"))
    conv._soffice = "soffice"
    src = tmp_path / "gammalt_dokument.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)

    assert conv.can_convert(str(src))
    md = conv.convert(str(src), "http://x/gammalt.doc", referer_url="http://x/sida")
    assert md and calls and "--convert-to" in calls[0] and "docx" in calls[0]
    text = open(md, encoding="utf-8").read()
    assert "Innehåll som LibreOffice konverterat" in text
    assert "Word (äldre)" in text
    assert "**Källa:** http://x/sida" in text

    conv._soffice = None
    assert not conv.can_convert(str(src))


def test_scanned_pdf_uses_ocr_when_available_else_reports(tmp_path, monkeypatch):
    import pymupdf
    pdf = tmp_path / "skannad.pdf"
    doc = pymupdf.open()
    doc.new_page()               # sida utan textlager = skannad bild
    doc.save(pdf)
    doc.close()

    conv = uwc.DocumentConverter(str(tmp_path / "out"))
    monkeypatch.setattr(conv, "ocr_available", lambda: False)
    assert conv.convert(str(pdf), "http://x/skannad.pdf") is None
    assert "Tesseract" in conv.last_error

    monkeypatch.setattr(conv, "ocr_available", lambda: True)
    monkeypatch.setattr(conv, "_ocr_page", lambda page: "Text som OCR läste ur den skannade sidan.")
    md = conv.convert(str(pdf), "http://x/skannad.pdf", referer_url="http://x/sida")
    assert md
    text = open(md, encoding="utf-8").read()
    assert "Text som OCR läste" in text
    assert "**Textkälla:** OCR" in text


# ───────────────────────── innehållskvalitet ─────────────────────────
def test_canonical_copy_is_not_saved_but_target_is_crawled(tmp_path):
    async def go():
        site = Site()
        site.pages["/"] = (page("Start", "start", [("/print/p", "Utskrift")]), None)
        body = page("Avgifter", "avgiftsinformation")
        site.pages["/print/p"] = (body.replace("</head>", "<link rel='canonical' href='/p'></head>"), None)
        site.pages["/p"] = (body, None)
        # felkonfigurerad: pekar på startsidan → ska INTE räknas som dubblett
        site.pages["/fel"] = (page("Fel", "x").replace("</head>", "<link rel='canonical' href='/'></head>"), None)
        site.pages["/"] = (page("Start", "start", [("/print/p", "u"), ("/fel", "f")]), None)
        runner, base = await serve(site)
        try:
            return await run_crawl(base, tmp_path)
        finally:
            await runner.cleanup()
    c = run(go())
    saved = texts(tmp_path)
    assert len(saved) == 3                      # start, /p (via canonical) och /fel
    assert len(c.report["canonical_skipped"]) == 1
    assert c.report["canonical_skipped"][0]["url"].endswith("/print/p")


def test_canonical_cycle_keeps_one_copy(tmp_path):
    async def go():
        site = Site()
        a = page("Samma", "x").replace("</head>", "<link rel='canonical' href='/c2'></head>")
        b = page("Samma", "x").replace("</head>", "<link rel='canonical' href='/c1'></head>")
        site.pages["/"] = (page("Start", "s", [("/c1", "1")]), None)
        site.pages["/c1"] = (a, None)
        site.pages["/c2"] = (b, None)
        runner, base = await serve(site)
        try:
            return await run_crawl(base, tmp_path)
        finally:
            await runner.cleanup()
    run(go())
    assert len(texts(tmp_path)) == 2            # start + exakt en av c1/c2


def test_identical_content_is_deduplicated_and_heals(tmp_path):
    async def go():
        site = Site()
        same = page("Samma sida", "identiskt innehåll")
        site.pages["/"] = (page("Start", "s", [("/a", "a"), ("/b", "b")]), None)
        site.pages["/a"] = (same, None)
        site.pages["/b"] = (same, None)
        runner, base = await serve(site)
        try:
            c1 = await run_crawl(base, tmp_path)
            n1 = len(texts(tmp_path))
            c2 = await run_crawl(base, tmp_path)            # omkörning: stabilt, inga raderingar
            n2 = len(texts(tmp_path))
            removed_2 = [c for c in read_changes(tmp_path) if c["event"] == "removed"]

            # Originalet försvinner → kopian ska sparas inom två körningar
            kept = next(c["url"] for c in read_changes(tmp_path)
                        if c["event"] == "added" and c["url"].endswith(("/a", "/b")))
            site.gone.add("/" + kept.rsplit("/", 1)[1])
            await run_crawl(base, tmp_path)
            await run_crawl(base, tmp_path)
        finally:
            await runner.cleanup()
        return c1, n1, n2, removed_2
    c1, n1, n2, removed_2 = run(go())
    assert n1 == 2 and len(c1.report["duplicates"]) == 1
    assert n2 == 2 and removed_2 == []
    assert "identiskt innehåll" in all_md(tmp_path)


def test_language_filter_skips_other_languages(tmp_path):
    async def go():
        site = Site()
        site.pages["/"] = (page("Start", "s", [("/en", "English"), ("/sv", "Svenska")]), None)
        site.pages["/en"] = (page("English page", "hello").replace("lang='sv'", "lang='en-GB'"), None)
        site.pages["/sv"] = (page("Svensk sida", "hej"), None)
        runner, base = await serve(site)
        try:
            return await run_crawl(base, tmp_path, languages=["sv"])
        finally:
            await runner.cleanup()
    c = run(go())
    assert len(texts(tmp_path)) == 2
    assert [w["language"] for w in c.report["wrong_language"]] == ["en"]


def test_sitemap_lastmod_skips_unchanged_pages(tmp_path):
    async def go():
        site = Site()
        site.pages["/"] = (page("Start", "s", [("/a", "a"), ("/b", "b")]), None)
        site.pages["/a"] = (page("Sida A", "a"), None)
        site.pages["/b"] = (page("Sida B", "b"), None)
        site.sitemap_paths = ["/a", "/b"]
        site.sitemap_lastmod = {"/a": "2020-01-01", "/b": "2999-01-01T00:00:00Z"}
        runner, base = await serve(site)
        try:
            await run_crawl(base, tmp_path)
            site.hits.clear()
            c2 = await run_crawl(base, tmp_path)
            hit_paths_2 = {h[0] for h in site.hits}
            site.hits.clear()
            await run_crawl(base, tmp_path, sitemap_lastmod_max_age_days=0)
            hit_paths_3 = {h[0] for h in site.hits}
        finally:
            await runner.cleanup()
        return c2, hit_paths_2, hit_paths_3
    c2, hits2, hits3 = run(go())
    assert "/a" not in hits2 and "/b" in hits2           # /a har gammal lastmod → ingen request
    assert c2.stats.pages_skipped_lastmod == 1
    assert c2.stats.pages_visited == 3                   # undersidor räknas ändå (länkar spelas upp)
    assert "/a" in hits3                                 # säkerhetsventilen: gamla poster kontrolleras igen


def test_parse_lastmod_variants():
    p = uwc.AsyncWebCrawler._parse_lastmod
    assert p("2024-05-01").day == 2                      # bara datum → slutet av dagen
    assert p("2024-05-01T10:00:00+02:00") is not None
    assert p("2024-05-01T10:00:00Z") is not None
    assert p("skräp") is None


def test_chunks_jsonl_covers_pages_and_documents(tmp_path):
    async def go():
        site = Site()
        site.pages["/"] = (page("Start", "s", [("/d/policy.pdf", "Policy för resor"),
                                               ("/a", "A")]), '"s"')
        site.pages["/a"] = (page("Sida A", "unik-text-a"), '"a"')
        site.files["/d/policy.pdf"] = (make_pdf("Reseriktlinjer i dokumentet"), '"p"')
        runner, base = await serve(site)
        try:
            await run_crawl(base, tmp_path)
            first = open(tmp_path / "chunks.jsonl", encoding="utf-8").read()
            await run_crawl(base, tmp_path)           # inkrementell körning ger samma innehåll
            second = open(tmp_path / "chunks.jsonl", encoding="utf-8").read()
        finally:
            await runner.cleanup()
        return first, second
    first, second = run(go())
    recs = [json.loads(l) for l in first.splitlines()]
    assert len({r["id"] for r in recs}) == len(recs)
    pages = [r for r in recs if r["source_type"] == "page"]
    docs = [r for r in recs if r["source_type"] == "document"]
    assert any("unik-text-a" in r["content"] for r in pages)
    assert docs and docs[0]["url"].endswith("/d/policy.pdf")
    assert docs[0]["referer_url"].startswith("http://127.0.0.1")
    assert all(r["url"] and r["context"] for r in recs)
    assert sorted(first.splitlines()) == sorted(second.splitlines())


def test_chunks_jsonl_from_json_format(tmp_path):
    async def go():
        site = Site()
        site.pages["/"] = (page("Start", "s"), None)
        runner, base = await serve(site)
        try:
            await run_crawl(base, tmp_path, save_format=".json")
        finally:
            await runner.cleanup()
    run(go())
    recs = [json.loads(l) for l in open(tmp_path / "chunks.jsonl", encoding="utf-8")]
    assert recs and recs[0]["source_type"] == "page" and recs[0]["heading_path"]
