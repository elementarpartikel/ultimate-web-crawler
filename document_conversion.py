"""PDF/Office/OCR utan GUI, crawler, HTML-parser eller webbläsarimport."""
from __future__ import annotations
import hashlib
import importlib.util
import math
import os
import re
import shutil
import subprocess
import tempfile
import threading
from collections import Counter
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from crawler_runtime import remove_tree
from crawler_text import (LogLevel, slugify, metadata_line, file_sha256, markdown_table,
                          downgrade_body_h1, atomic_write_text, detect_prompt_injection, clean_pii_text,
                          _split_long_text)

HAS_PYMUPDF = bool(importlib.util.find_spec('pymupdf') or importlib.util.find_spec('fitz'))
HAS_DOCX = bool(importlib.util.find_spec('docx'))
HAS_OPENPYXL = bool(importlib.util.find_spec('openpyxl'))
HAS_PPTX = bool(importlib.util.find_spec('pptx'))


def __getattr__(name):
    """En PDF-worker ska inte importera Office-bibliotek i onödan."""
    # Statiska importuttryck gör biblioteken synliga även för paketeringsverktyg.
    if name == 'pymupdf':
        try:
            import pymupdf as value
        except ImportError:
            import fitz as value
    elif name == 'DocxDocument':
        from docx import Document as value
    elif name == 'DocxTable':
        from docx.table import Table as value
    elif name == 'DocxParagraph':
        from docx.text.paragraph import Paragraph as value
    elif name == 'xl_load_workbook':
        from openpyxl import load_workbook as value
    elif name == 'PptxPresentation':
        from pptx import Presentation as value
    else:
        raise AttributeError(name)
    globals()[name] = value
    return value


def _load_parser(ext):
    if ext == '.pdf':
        if 'pymupdf' not in globals():
            __getattr__('pymupdf')
    elif ext in ('.docx', '.dotx'):
        for name in ('DocxDocument', 'DocxTable', 'DocxParagraph'):
            if name not in globals():
                __getattr__(name)
    elif ext == '.xlsx' and 'xl_load_workbook' not in globals():
        __getattr__('xl_load_workbook')
    elif ext == '.pptx' and 'PptxPresentation' not in globals():
        __getattr__('PptxPresentation')

class DocumentConverter:
    """Konverterar binära dokument (PDF, Word, Excel, PPTX) till Markdown.

    Extraherar text och skapar en .md-fil med intranät-URL:en i toppen,
    så att RAG-system kan citera rätt källa även för dokument-text.
    Originaldokumentet behålls alltid i dokument/-mappen.
    """

    SUPPORTED = {
        '.pdf': 'PDF', '.docx': 'Word', '.dotx': 'Word-mall',
        '.xlsx': 'Excel', '.pptx': 'PowerPoint',
    }
    # Äldre format stöds inte direkt av Python-biblioteken. Om LibreOffice
    # (soffice) finns på datorn konverteras de först till moderna format.
    LEGACY_TARGETS = {
        '.doc': 'docx', '.rtf': 'docx', '.odt': 'docx',
        '.xls': 'xlsx', '.ods': 'xlsx',
        '.ppt': 'pptx', '.odp': 'pptx',
    }
    LEGACY_NAMES = {
        '.doc': 'Word (äldre)', '.rtf': 'RTF', '.odt': 'OpenDocument Text',
        '.xls': 'Excel (äldre)', '.ods': 'OpenDocument Calc',
        '.ppt': 'PowerPoint (äldre)', '.odp': 'OpenDocument Presentation',
    }
    MAX_XLSX_ROWS_PER_SHEET = 5000
    LEGACY_TIMEOUT_S = 120

    # Metadata-titlar som är skräp och inte ska användas som dokumenttitel.
    _JUNK_TITLE_RE = re.compile(
        r'^(microsoft (word|excel|powerpoint)\b.*|untitled.*|namnlöst.*|'
        r'document\s*\d*|dokument\s*\d*|powerpoint presentation|presentation\d*|'
        r'slide 1|bild 1|.*\.(docx?|xlsx?|pptx?|pdf|odt)|\s*)$',
        re.IGNORECASE)

    # Generiska länktexter som inte är beskrivande nog att använda som
    # dokument-titel. Matchning sker case-insensitivt efter strip().
    GENERIC_LINK_TEXTS = {
        'ladda ner', 'ladda ner fil', 'ladda ned', 'ladda ned fil',
        'hämta', 'hämta fil', 'download', 'download file',
        'öppna', 'öppna fil', 'öppna dokument',
        'visa', 'visa fil', 'visa dokument',
        'klicka här', 'läs mer', 'read more', 'click here',
        'länk', 'link', 'pdf', 'dokument', 'document',
        '(via sitemap)', 'via sitemap',
    }

    def __init__(self, output_dir: str,
                 log_fn=None, pii_cleaner=None, ocr: bool = True,
                 ocr_languages: str = "swe+eng", ocr_max_pages: int = 50,
                 work_dir: Optional[str] = None):
        self.texts_dir = os.path.join(output_dir, "texter")
        self._log = log_fn
        self._clean_pii = pii_cleaner
        self.ocr_enabled = ocr
        self.ocr_languages = ocr_languages
        self.ocr_max_pages = ocr_max_pages
        self._ocr_ok: Optional[bool] = None
        self.work_dir = work_dir
        self.last_used_ocr = False    # True om senaste dokumentets text kommer från OCR
        self.last_error = ""          # orsak till senaste misslyckade konvertering
        self.last_raw_injection = False
        self._soffice = (shutil.which("soffice") or shutil.which("libreoffice")
                         or shutil.which("soffice.exe"))
        os.makedirs(self.texts_dir, exist_ok=True)

    @staticmethod
    def md_filename(doc_url: str, filename: str) -> str:
        """Stabilt, kollisionsfritt namn på den konverterade .md-filen.

        Bygger alltid på samma två delar — en kapad slug av filnamnet och en
        hash av dokument-URL:en — så att crawlern och konverteraren aldrig kan
        komma överens om olika namn, och så att hashen aldrig kapas bort.
        """
        base = os.path.splitext(os.path.basename(filename))[0]
        base = re.sub(r'_[0-9a-f]{6}$', '', base)          # ta bort ev. gammal hash
        slug = slugify(base)[:40] or "dokument"
        url_hash = hashlib.md5(doc_url.encode('utf-8')).hexdigest()[:8]
        return f"{slug}_{url_hash}_doc.md"

    def md_path_for(self, doc_url: str, filename: str) -> str:
        return os.path.join(self.texts_dir, self.md_filename(doc_url, filename))

    def can_convert(self, filepath: str) -> bool:
        ext = os.path.splitext(filepath)[1].lower()
        if ext in self.LEGACY_TARGETS:
            return bool(self._soffice)
        if ext == '.pdf' and not HAS_PYMUPDF:
            return False
        if ext in ('.docx', '.dotx') and not HAS_DOCX:
            return False
        if ext == '.xlsx' and not HAS_OPENPYXL:
            return False
        if ext == '.pptx' and not HAS_PPTX:
            return False
        return ext in self.SUPPORTED

    def _legacy_to_modern(self, filepath: str, ext: str, workdir: str) -> Optional[str]:
        """Konverterar .doc/.xls/.ppt m.fl. med LibreOffice. Returnerar sökväg eller None."""
        target = self.LEGACY_TARGETS[ext]
        try:
            subprocess.run(
                [self._soffice, '-env:UserInstallation=' + Path(
                    os.path.join(workdir, 'lo-profile')).as_uri(),
                 "--headless", "--norestore", "--convert-to", target,
                 "--outdir", workdir, filepath],
                check=True, capture_output=True, timeout=self.LEGACY_TIMEOUT_S)
        except Exception as e:
            self.last_error = f"LibreOffice-konvertering misslyckades: {e}"
            return None
        out = os.path.join(workdir, os.path.splitext(os.path.basename(filepath))[0]
                           + "." + target)
        return out if os.path.exists(out) else None

    def convert(self, filepath: str, doc_url: str,
                referer_url: str = "", referer_title: str = "",
                link_text: str = "", fetched_at: Optional[str] = None,
                server_last_modified: str = '') -> Optional[str]:
        """Konverterar dokument till .md. Returnerar sökväg eller None.

        Metadata bäddas in i brödtexten (inte bara i header) eftersom RAG-
        pipelines ofta kapar de första raderna före retrieval. Källan
        upprepas också sist i filen — om Sveas chunking splittrar dokumentet
        i flera bitar säkrar det att åtminstone en chunk har källan med.
        """
        self.last_error = ""
        self.last_used_ocr = False
        self.last_raw_injection = False
        tmpdir = None
        ext = os.path.splitext(filepath)[1].lower()
        src_path, src_ext = filepath, ext
        try:
            if ext in self.LEGACY_TARGETS:
                if not self._soffice:
                    self.last_error = "Äldre format kräver LibreOffice (soffice)"
                    return None
                tmpdir = tempfile.mkdtemp(prefix="uwc_", dir=self.work_dir)
                src_path = self._legacy_to_modern(filepath, ext, tmpdir)
                if not src_path:
                    return None
                src_ext = os.path.splitext(src_path)[1].lower()

            extractors = {
                '.pdf': self._extract_pdf,
                '.docx': self._extract_docx,
                '.dotx': self._extract_docx,
                '.xlsx': self._extract_xlsx,
                '.pptx': self._extract_pptx,
            }
            extractor = extractors.get(src_ext)
            if not extractor:
                self.last_error = f"Filtypen {ext} stöds inte"
                return None

            text = extractor(src_path)
            if not text or len(text.strip()) < 20:
                if ext == '.pdf':
                    if self.ocr_available():
                        self.last_error = "Ingen text i PDF:en, och OCR gav heller ingen text"
                    else:
                        self.last_error = ("Ingen text i PDF:en — troligen en skannad bild. "
                                           "Installera Tesseract för OCR")
                else:
                    self.last_error = "För lite text i dokumentet"
                if self._log:
                    self._log(
                        f"  ⚠ Konvertering gav för lite text "
                        f"({len(text.strip()) if text else 0} tecken): "
                        f"{os.path.basename(filepath)} — {self.last_error}",
                        LogLevel.WARNING)
                return None

            if self._clean_pii:
                text = self._clean_pii(text)

            # Titel-prioritet (faller från bäst till sista utvägen):
            #   1. link_text   — <a>-taggens text (t.ex. "Eko ack 2025 Januari"),
            #                    kräver >5 tecken OCH att texten inte är generisk
            #                    ("Ladda ner fil", "Download", etc.)
            #   2. metadata_title — dokumentets egen title (PDF/Office metadata)
            #   3. humaniserat filnamn ("Eko ack 2025 januari")
            #   4. referer_title — sidans <title> (ofta för generisk,
            #                      t.ex. "Startsida - Tyresö kommun")
            #   5. rå slug som sista utväg
            filename = os.path.basename(filepath)
            base_name = os.path.splitext(filename)[0]
            metadata_title = self._extract_doc_metadata_title(src_path, src_ext)
            if metadata_title and self._JUNK_TITLE_RE.match(metadata_title.strip()):
                metadata_title = ""
            humanized = self._humanize_filename(base_name)
            clean_link = (link_text or "").strip()

            # De-duplicera dubbla länktexter ("Ladda ner fil Ladda ner fil")
            if clean_link:
                total_len = len(clean_link)
                if total_len % 2 == 1:  # udda längd → testa med mellanslag i mitten
                    mid = total_len // 2
                    if (clean_link[mid] == ' '
                            and clean_link[:mid] == clean_link[mid + 1:]):
                        clean_link = clean_link[:mid]
                elif total_len % 2 == 0:
                    mid = total_len // 2
                    if clean_link[:mid].rstrip() == clean_link[mid:].lstrip():
                        clean_link = clean_link[:mid].rstrip()

            def _is_usable_link_text(text: str) -> bool:
                """True om link_text är tillräckligt beskrivande för titel."""
                if len(text) <= 5:
                    return False
                return text.lower() not in self.GENERIC_LINK_TEXTS

            raw_title = (
                (clean_link if _is_usable_link_text(clean_link) else "")
                or metadata_title
                or humanized
                or (referer_title.strip() if referer_title else "")
                or base_name
            )

            # Strippa storleks-suffix som "pdf, 535.8 kB, öppnas i nytt fönster"
            title = re.sub(
                r'\s*(pdf|docx?|xlsx?|pptx?|pages?)?\s*,?\s*'
                r'\d+(?:[.,]\d+)?\s*[kKmMgG]?[bB][, .].*$',
                '', raw_title, flags=re.IGNORECASE).strip()
            title = re.sub(r'\s*[, .]+\s*$', '', title) or raw_title
            title = metadata_line(title)
            # Titeln kan komma från länktext/metadata som inte passerat PII-tvätten
            if self._clean_pii:
                title = self._clean_pii(title)
                if self._clean_pii and referer_title:
                    referer_title = self._clean_pii(referer_title)

            file_type = (self.SUPPORTED.get(ext) or self.LEGACY_NAMES.get(ext)
                         or "Dokument")

            # Bygg metadata-block som SYNLIG brödtext (fetstil), inte ren header.
            # Sveas chunker skippar typiskt header-text men behåller brödtext.
            meta_lines = [f"# {title}", ""]
            meta_lines.append(f"**Källa:** {referer_url or doc_url}")
            if doc_url:
                meta_lines.append(f"**Dokument-URL:** {doc_url}")
            if filename:
                # Visa originalfilnamnet (utan vår interna hash-suffix)
                display_name = re.sub(r'_[0-9a-f]{6}(\.\w+)$', r'\1', filename)
                display_name = metadata_line(self._clean_pii(display_name) if self._clean_pii else display_name)
                meta_lines.append(f"**Filnamn:** {display_name}")
            meta_lines.append(f"**Filtyp:** {file_type}")
            # LibreOffice kan ändra temporärfilens metadata under konverteringen.
            dates = self._extract_doc_dates(src_path, src_ext) if src_path == filepath else {}
            fetched = datetime.now().isoformat(timespec='seconds') if fetched_at is None else fetched_at
            server_modified = self._parse_date(server_last_modified, http=True)
            date_lines = []
            for label, value in (('Hämtad', fetched), ('Dokumentets ändringsdatum', dates.get('modified')),
                                 ('Serverns ändringsdatum', server_modified), ('Skapad', dates.get('created'))):
                if value:
                    date_lines.append(f'**{label}:** {metadata_line(str(value))}')
            meta_lines.extend(date_lines)
            if self.last_raw_injection or detect_prompt_injection(text, title):
                meta_lines.append('**Flaggor:** possible_prompt_injection')
            if self.last_used_ocr:
                meta_lines.append("**Textkälla:** OCR (automatisk textigenkänning — kan innehålla fel)")
            meta_lines.append("")
            meta_lines.append("---")
            meta_lines.append("")

            # Upprepa källan sist i filen — om dokumentet chunkas i flera
            # bitar har åtminstone första och sista bitarna källinformation.
            footer_lines = ["", "", "---", ""]
            footer_lines.append(f"**Källa:** {referer_url or doc_url}")
            if doc_url:
                footer_lines.append(f"**Dokument-URL:** {doc_url}")

            text = downgrade_body_h1(text)
            # Gör varje sida/bild och långa avsnitt begripliga också vid extern chunkning.
            text = self._repeat_document_context(text, title, doc_url, referer_url, date_lines)
            md_content = ("\n".join(meta_lines) + text
                          + "\n".join(footer_lines))

            md_path = self.md_path_for(doc_url, filename)

            atomic_write_text(md_path, md_content)
            return md_path

        except Exception as e:
            self.last_error = f"Konvertering misslyckades: {e}"
            if self._log:
                self._log(
                    f"  ✗ Konvertering misslyckades "
                    f"({os.path.basename(filepath)}): {e}", LogLevel.ERROR)
            return None
        finally:
            if tmpdir:
                remove_tree(self.work_dir or tempfile.gettempdir(), tmpdir)

    # ─── Titelextraktion ────────────────────────────────────
    @staticmethod
    def _repeat_document_context(text: str, title: str, doc_url: str,
                                  referer_url: str, dates: List[str]) -> str:
        context = '\n'.join([f'**Dokument:** {title}', f'**Källa:** {referer_url or doc_url}',
                             f'**Dokument-URL:** {doc_url}', *dates])
        result = []
        for part in re.split(r'(?=^##\s)', text, flags=re.MULTILINE):
            part = part.strip()
            if not part:
                continue
            match = re.match(r'^##\s+([^\n]+)\n?', part)
            heading = match.group(1).strip() if match else title
            body = part[match.end():].strip() if match else part
            unit = re.match(r'^(Sida|Bild)\s+(\d+)\b', heading, re.IGNORECASE)
            if unit:
                heading = f'{title} – {unit.group(1).lower()} {unit.group(2)}'
            for piece in _split_long_text(body, 400, 0) if body else ['']:
                result.append(f'## {heading}\n\n{context}\n\n{piece}')
        return '\n\n'.join(result)

    @staticmethod
    def _parse_date(value, http: bool = False) -> str:
        if not value:
            return ''
        try:
            if isinstance(value, datetime):
                parsed = value
            elif http:
                parsed = parsedate_to_datetime(str(value))
            elif str(value).startswith('D:'):
                match = re.fullmatch(r"D:(\d{4})(\d{2})(\d{2})(\d{2})?(\d{2})?(\d{2})?(Z|[+-]\d{2}'?\d{2}'?)?", str(value))
                if not match:
                    return ''
                year, month, day, hour, minute, second, offset = match.groups()
                parsed = datetime(int(year), int(month), int(day), int(hour or 0),
                                  int(minute or 0), int(second or 0))
                if offset:
                    cleaned = offset.replace("'", '')
                    if cleaned == 'Z':
                        parsed = parsed.replace(tzinfo=timezone.utc)
                    else:
                        sign = 1 if cleaned[0] == '+' else -1
                        parsed = parsed.replace(tzinfo=timezone(sign * timedelta(
                            hours=int(cleaned[1:3]), minutes=int(cleaned[3:5]))))
                if hour is None:
                    return parsed.date().isoformat()
            else:
                parsed = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
            return parsed.isoformat(timespec='seconds')
        except (ValueError, TypeError, OverflowError):
            return ''

    def _extract_doc_dates(self, filepath: str, ext: str) -> Dict:
        # Metadata beskriver filens tekniska historik, aldrig taxans giltighet.
        try:
            _load_parser(ext)
            if ext == '.pdf' and HAS_PYMUPDF:
                with pymupdf.open(filepath) as doc:
                    meta = doc.metadata or {}
                    return {'created': self._parse_date(meta.get('creationDate')),
                            'modified': self._parse_date(meta.get('modDate'))}
            if ext in ('.docx', '.dotx'):
                props = DocxDocument(filepath).core_properties
            elif ext == '.pptx':
                props = PptxPresentation(filepath).core_properties
            elif ext == '.xlsx':
                wb = xl_load_workbook(filepath, read_only=True, data_only=True)
                try:
                    return {'created': self._parse_date(wb.properties.created),
                            'modified': self._parse_date(wb.properties.modified)}
                finally:
                    wb.close()
            else:
                return {}
            return {'created': self._parse_date(props.created), 'modified': self._parse_date(props.modified)}
        except Exception:
            return {}

    @staticmethod
    def _humanize_filename(name: str) -> str:
        """Gör ett slug-namn läsbart: 'ansokan-om-x_ad7919' → 'Ansokan om x'.

        Eftersom slugify har strippat svenska tecken (ä→a, ö→o) går det
        inte att återskapa originalnamnet helt — men "Ansokan om
        parkeringstillstand" är vida bättre än "ansokan-om-parkeringstillstand_ad7919".
        """
        # Ta bort hash-suffix
        name = re.sub(r'_[0-9a-f]{6}$', '', name)
        # Bindestreck och underscore → mellanslag
        name = re.sub(r'[-_]+', ' ', name).strip()
        # Stor första bokstav (ev. övriga ord lämnas som de är —
        # vi vet inte vilka som är egennamn)
        if name:
            name = name[0].upper() + name[1:]
        return name

    def _extract_doc_metadata_title(self, filepath: str, ext: str) -> str:
        """Plocka ut dokumenttitel ur filens egen metadata om sådan finns.

        Många PDF/Office-dokument har ett `title`-fält i sin metadata som
        ofta är mer beskrivande än filnamnet ("Ansökan om parkeringstillstånd
        för rörelsehindrad" vs "ansokan-om-parkeringstillstand").
        """
        try:
            _load_parser(ext)
            if ext == '.pdf' and HAS_PYMUPDF:
                doc = pymupdf.open(filepath)
                title = (doc.metadata or {}).get('title', '')
                doc.close()
                return title.strip() if title else ""
            if ext in ('.docx', '.dotx') and HAS_DOCX:
                doc = DocxDocument(filepath)
                title = (doc.core_properties.title or '').strip()
                return title
            if ext == '.xlsx' and HAS_OPENPYXL:
                wb = xl_load_workbook(filepath, read_only=True, data_only=True)
                title = (wb.properties.title or '').strip()
                wb.close()
                return title
            if ext == '.pptx' and HAS_PPTX:
                prs = PptxPresentation(filepath)
                title = (prs.core_properties.title or '').strip()
                return title
        except Exception:
            pass
        return ""

    def ocr_available(self) -> bool:
        """OCR kräver att Tesseract är installerat (PyMuPDF anropar det)."""
        if not (HAS_PYMUPDF and self.ocr_enabled):
            return False
        if self._ocr_ok is None:
            try:
                _load_parser('.pdf')
                pymupdf.get_tessdata()
                self._ocr_ok = True
            except Exception:
                self._ocr_ok = False
        return self._ocr_ok

    def _ocr_page(self, page) -> str:
        """OCR på en sida. Provar önskade språk, sen enbart engelska."""
        for lang in (self.ocr_languages, "eng"):
            try:
                textpage = page.get_textpage_ocr(language=lang, dpi=200, full=True)
                return page.get_text("text", textpage=textpage)
            except Exception:
                continue
        return ""

    def _extract_pdf(self, filepath: str) -> str:
        _load_parser('.pdf')
        doc = pymupdf.open(filepath)
        pages, font_weights = [], Counter()
        margin_counts = Counter()
        ocr_attempts = 0
        try:
            for i, page in enumerate(doc):
                text = page.get_text('text', sort=True)
                self.last_raw_injection |= detect_prompt_injection(text)
                if not text.strip() and ocr_attempts < self.ocr_max_pages and self.ocr_available():
                    # Sidan saknar textlager (skannad bild) → OCR
                    ocr_attempts += 1
                    text = self._ocr_page(page)
                    if text.strip():
                        self.last_used_ocr = True
                        self.last_raw_injection |= detect_prompt_injection(text)
                    pages.append({'number': i + 1, 'items': [], 'fallback': text.strip(),
                                  'width': page.rect.width})
                    continue
                items, covered = [], []
                try:
                    tables = page.find_tables().tables
                    for table in tables:
                        rows = table.extract()
                        if not rows or table.col_count < 2 or table.row_count < 2:
                            continue
                        header = table.header.names if table.header.external else None
                        table_text = markdown_table(rows, header)
                        if not table_text:
                            continue
                        bbox = tuple(table.bbox)
                        covered.append(bbox)
                        if table.header.external:
                            covered.append(tuple(table.header.bbox))
                        items.append({'text': table_text, 'bbox': bbox, 'table': True})
                except Exception:
                    # Misslyckad tabellidentifiering får inte ta bort råtexten.
                    covered = []
                    items = []
                try:
                    blocks = page.get_text('dict', sort=True)['blocks']
                    for block in blocks:
                        for line in block.get('lines', []):
                            spans = line.get('spans', [])
                            line_text = self._pdf_line_text(spans)
                            if not line_text:
                                continue
                            bbox = tuple(line['bbox'])
                            cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
                            if any(x0 <= cx <= x1 and y0 <= cy <= y1 for x0, y0, x1, y1 in covered):
                                continue
                            size = max((span.get('size', 0) for span in spans), default=0)
                            bold = all(span.get('flags', 0) & 16 for span in spans if span.get('text', '').strip())
                            margin = ('top' if bbox[3] < page.rect.height * 0.1 else
                                      'bottom' if bbox[1] > page.rect.height * 0.9 else '')
                            key = self._pdf_margin_key(line_text, margin)
                            items.append({'text': line_text, 'bbox': bbox, 'size': size,
                                          'bold': bold, 'margin_key': key})
                            if not margin:
                                for span in spans:
                                    font_weights[round(span.get('size', 0), 1)] += len(span.get('text', ''))
                    margin_counts.update({item['margin_key'] for item in items if item.get('margin_key')})
                except Exception:
                    items = []
                pages.append({'number': i + 1, 'items': items, 'fallback': text.strip(),
                              'width': page.rect.width})
        finally:
            doc.close()
        body_size = font_weights.most_common(1)[0][0] if font_weights else 11
        repeated = {key for key, count in margin_counts.items()
                    if len(pages) >= 3 and count >= max(3, math.ceil(len(pages) * 0.6))}
        output = []
        for record in pages:
            items = [item for item in record['items'] if item.get('margin_key') not in repeated]
            lines = []
            for item in self._pdf_reading_order(items, record['width']):
                value = item['text']
                size = item.get('size', 0)
                if (not item.get('table') and not item.get('margin_key') and len(value) < 160
                        and (size >= body_size * 1.2 or (item.get('bold') and size >= body_size))):
                    level = 3 if size >= body_size * 1.4 else 4
                    value = '#' * level + ' ' + value
                lines.append(value)
            # Om all text är en upprepad sidfot ska den inte återinföras via fallback.
            value = '\n\n'.join(lines) if record['items'] else record['fallback']
            if value.strip():
                output.append(f"## Sida {record['number']}\n\n{value.strip()}")
        return '\n\n'.join(output)

    @staticmethod
    def _pdf_line_text(spans: List[Dict]) -> str:
        value, previous = '', None
        for span in spans:
            part = span.get('text', '')
            if (previous and value and part and not value[-1].isspace() and not part[0].isspace()
                    and span['bbox'][0] - previous['bbox'][2] > span.get('size', 10) * 0.15):
                value += ' '
            value += part
            previous = span
        return value.strip()

    @staticmethod
    def _pdf_margin_key(text: str, margin: str) -> str:
        if not margin or len(text) > 200:
            return ''
        key = re.sub(r'\s+', ' ', text).strip().lower()
        key = re.sub(r'\b(sida|page)\s+\d+(?:\s+(av|of)\s+\d+)?', r'\1 #', key)
        if re.fullmatch(r'\d+(?:\s*/\s*\d+)?', key):
            key = '#'
        return margin + ':' + key

    @staticmethod
    def _pdf_reading_order(items: List[Dict], width: float) -> List[Dict]:
        """Läs tydliga tvåspaltsband spaltvis, med breda rubriker/tabeller emellan."""
        mid = width / 2
        left = [item for item in items if item['bbox'][2] < mid and not item.get('margin_key')]
        right = [item for item in items if item['bbox'][0] > mid and not item.get('margin_key')]
        normal = lambda values: sorted(values, key=lambda item: (item['bbox'][1], item['bbox'][0]))
        if len(left) < 3 or len(right) < 3:
            return normal(items)
        # Priser/etiketter på samma rad är en tabell utan linjer, inte två
        # självständiga textspalter. Bevara kopplingen också i chunkad text.
        pairs, unused = [], list(right)
        for label in normal(left):
            candidates = [value for value in unused
                          if abs(value['bbox'][3] - label['bbox'][3]) <= max(
                              3, (label['bbox'][3] - label['bbox'][1]) * 0.3)]
            if candidates:
                value = min(candidates, key=lambda item: abs(item['bbox'][3] - label['bbox'][3]))
                pairs.append((label, value))
                unused.remove(value)
        def tabular(label, value):
            numeric = bool(re.fullmatch(r'[\d\s.,/%€$:+−–-]+(?:\s*(?:kr|sek|eur|m²|m2|st|kg))?',
                                        value['text'].strip(), re.I))
            status = bool(re.fullmatch(r'ja|nej|yes|no', value['text'].strip(), re.I))
            short = len(label['text'].split()) <= 4 and len(value['text'].split()) <= 3
            return numeric or status or short
        table_pairs = [(label, value) for label, value in pairs if tabular(label, value)]
        if (len(pairs) >= min(len(left), len(right)) * 0.7 and pairs
                and len(table_pairs) >= 3):
            paired = {id(item) for pair in table_pairs for item in pair}
            merged = [{**label, 'text': label['text'] + ' | ' + value['text'],
                       'bbox': (label['bbox'][0], min(label['bbox'][1], value['bbox'][1]),
                                value['bbox'][2], max(label['bbox'][3], value['bbox'][3])),
                       'table': True} for label, value in table_pairs]
            return DocumentConverter._pdf_reading_order(
                merged + [item for item in items if id(item) not in paired], width)
        if max(min(item['bbox'][1] for item in left), min(item['bbox'][1] for item in right)) >= min(
                max(item['bbox'][3] for item in left), max(item['bbox'][3] for item in right)):
            return normal(items)
        broad = normal([item for item in items if item not in left and item not in right])
        remaining, ordered = list(left + right), []
        for separator in broad:
            band = [item for item in remaining if item['bbox'][1] < separator['bbox'][1]]
            ordered.extend(normal([item for item in band if item['bbox'][2] < mid]))
            ordered.extend(normal([item for item in band if item['bbox'][0] > mid]))
            remaining = [item for item in remaining if item not in band]
            ordered.append(separator)
        ordered.extend(normal([item for item in remaining if item['bbox'][2] < mid]))
        ordered.extend(normal([item for item in remaining if item['bbox'][0] > mid]))
        return ordered

    def _extract_docx(self, filepath: str) -> str:
        _load_parser('.docx')
        doc = DocxDocument(filepath)
        parts = []

        def para_to_md(para) -> Optional[str]:
            text = para.text.strip()
            if not text:
                return None
            style = para.style.name if para.style and para.style.name else ""
            if style.startswith('Heading') or style.startswith('Rubrik'):
                digits = re.findall(r'\d+', style)
                level = int(digits[0]) if digits else 1
                return f"{'#' * min(level + 1, 6)} {text}"
            if style == 'Title' or style == 'Titel':
                return f"## {text}"
            return text

        def table_to_md(table) -> Optional[str]:
            return markdown_table([[cell.text.strip() for cell in row.cells] for row in table.rows]) or None

        if DocxParagraph is not None and DocxTable is not None:
            # Gå igenom dokumentet i ordning så att tabeller hamnar där de står
            for child in doc.element.body.iterchildren():
                if child.tag.endswith('}p'):
                    md = para_to_md(DocxParagraph(child, doc))
                elif child.tag.endswith('}tbl'):
                    md = table_to_md(DocxTable(child, doc))
                else:
                    continue
                if md:
                    parts.append(md)
        else:
            for para in doc.paragraphs:
                md = para_to_md(para)
                if md:
                    parts.append(md)
            for table in doc.tables:
                md = table_to_md(table)
                if md:
                    parts.append(md)
        return "\n\n".join(parts)

    def _extract_xlsx(self, filepath: str) -> str:
        _load_parser('.xlsx')
        wb = xl_load_workbook(filepath, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue                      # hoppa över dolda blad
            parts.append(f"## {sheet_name}")
            rows_data = []
            truncated = False
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= self.MAX_XLSX_ROWS_PER_SHEET:
                    truncated = True
                    break
                cells = [str(c).replace('\n', ' ').replace('|', '/') if c is not None
                         else "" for c in row]
                if any(c for c in cells):
                    rows_data.append("| " + " | ".join(cells) + " |")
            if rows_data:
                col_count = rows_data[0].count("|") - 1
                rows_data.insert(1, "|" + "|".join(["---"] * max(col_count, 1)) + "|")
                parts.append("\n".join(rows_data))
            if truncated:
                parts.append(f"_(Bladet avkortat efter {self.MAX_XLSX_ROWS_PER_SHEET} rader.)_")
        wb.close()
        return "\n\n".join(parts)

    def _extract_pptx(self, filepath: str) -> str:
        _load_parser('.pptx')
        prs = PptxPresentation(filepath)
        parts = []

        def shape_text(shape, depth=0):
            if depth > 50:
                raise ValueError('För många nästlade PowerPoint-grupper')
            if hasattr(shape, 'shapes'):
                return [part for child in sorted(shape.shapes, key=lambda s: (s.top, s.left))
                        for part in shape_text(child, depth + 1)]
            if shape.has_table:
                return [markdown_table([[cell.text.strip() if not cell.is_spanned else ''
                                         for cell in row.cells] for row in shape.table.rows])]
            if shape.has_text_frame:
                return [para.text.strip() for para in shape.text_frame.paragraphs if para.text.strip()]
            return []

        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
                slide_texts.extend(shape_text(shape))
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = frame.text.strip() if frame is not None else ''
                if notes:
                    slide_texts.append('### Talarnoteringar\n\n' + notes)
            if slide_texts:
                parts.append(
                    f"## Bild {i + 1}\n\n" + "\n\n".join(slide_texts))
        return "\n\n".join(parts)

_DOCUMENT_CONVERSION_LOCK = threading.Lock()

def convert_document_result(converter, filepath: str, url: str, ref: Dict) -> Dict:
    """Fångar resultat och diagnostik under samma lås, utan delad resultatstatus."""
    with _DOCUMENT_CONVERSION_LOCK:
        path = converter.convert(filepath, url,
                                 referer_url=ref.get('referer_url', ''),
                                 referer_title=ref.get('referer_title', ''),
                                 link_text=ref.get('link_text', ''),
                                 **({key: ref[key] for key in ('fetched_at', 'server_last_modified') if key in ref}))
        return {"path": path, "ocr": converter.last_used_ocr,
                "error": converter.last_error}
